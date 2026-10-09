# -*- coding: utf-8 -*-
"""Excel appendix (all report tables, one sheet per chapter + an index sheet) from .cache/report_blocks.json.

Run: python scripts/report_build/build_xlsx.py     Output: reports/phase1_final_report_tables_appendix_v2.xlsx
"""
import io
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
BLOCKS = Path(__file__).resolve().parent / ".cache" / "report_blocks.json"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "reports" / "phase1_final_report_tables_appendix_v2.xlsx"

blocks = json.loads(io.open(BLOCKS, encoding="utf-8").read())
flat = lambda runs: "".join(r["text"] for r in runs)
EN = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")

sheets = {}          # sheet title -> list of (caption, header, rows)
order = []
cur_sheet, cur_caption, started = None, None, False
for b in blocks:
    if b["type"] == "heading" and b["level"] == 1:
        t = flat(b["runs"]).strip()
        m = re.match(r"^فصل\s+([۰-۹]+)", t)
        if m:
            cur_sheet = "فصل " + m.group(1)
            started = True
        elif t.startswith("خلاصه‌ی مدیریتی"):
            cur_sheet, started = "خلاصه‌ی مدیریتی", True
        else:
            cur_sheet = None
        cur_caption = None
    elif b["type"] == "para" and b["runs"] and b["runs"][0]["bold"] and flat(b["runs"]).strip().startswith("جدول"):
        cur_caption = flat(b["runs"]).strip().rstrip("*")
    elif b["type"] == "table" and started and cur_sheet:
        if cur_sheet not in sheets:
            sheets[cur_sheet] = []
            order.append(cur_sheet)
        sheets[cur_sheet].append((cur_caption or "جدول", [flat(c) for c in b["header"]], [[flat(c) for c in r] for r in b["rows"]]))
        cur_caption = None

wb = Workbook()
idx = wb.active
idx.title = "فهرست جداول"
idx.sheet_view.rightToLeft = True
head_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
cap_font = Font(bold=True, size=12, color="1F4E79")
al = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)

idx.append(["فصل", "عنوان جدول", "برگه"])
for c in idx[1]:
    c.font, c.fill, c.alignment = Font(bold=True), head_fill, al
total = 0
for title in order:
    ws = wb.create_sheet(title[:31])
    ws.sheet_view.rightToLeft = True
    r = 1
    widths = {}
    for cap, header, rows in sheets[title]:
        ws.cell(row=r, column=1, value=cap).font = cap_font
        row_idx = r
        r += 1
        for ci, h in enumerate(header, start=1):
            c = ws.cell(row=r, column=ci, value=h)
            c.font, c.fill, c.alignment = Font(bold=True), head_fill, al
            widths[ci] = max(widths.get(ci, 8), min(48, len(h) + 4))
        r += 1
        for row in rows:
            for ci, v in enumerate(row, start=1):
                v = v.replace("**", "")
                ws.cell(row=r, column=ci, value=v).alignment = al
                widths[ci] = max(widths.get(ci, 8), min(48, len(v) + 3))
            r += 1
        r += 2
        total += 1
        idx.append([title, cap, None])
        link = idx.cell(row=idx.max_row, column=3, value="باز کردن")
        link.hyperlink = f"#'{title[:31]}'!A{row_idx}"
        link.font = Font(color="0563C1", underline="single")
    for ci, w in widths.items():
        ws.column_dimensions[get_column_letter(ci)].width = w
idx.column_dimensions["A"].width = 18
idx.column_dimensions["B"].width = 80
idx.column_dimensions["C"].width = 14
for row in idx.iter_rows(min_row=2):
    for c in row[:2]:
        c.alignment = al

# --- پیوست: سری ماهانه‌ی نسبت مطالبات غیرجاری (بانک و نمونه) ---
import pandas as pd
_c = pd.read_csv(ROOT / "data" / "processed" / "phase1_market_overview" / "bank_coverage_monthly.csv")
_c = _c[_c.npl_coverage_reliable.astype(str) == "True"].copy()
_c["bank_ratio"] = 100 * _c.bank_npl / _c.bank_loans
_c["sample_ratio"] = 100 * _c.sample_npl / _c.sample_loans
_c["diff"] = _c.sample_ratio - _c.bank_ratio
ws2 = wb.create_sheet("پیوست سری NPL")
ws2.append(["ماه", "نسبت رسمی بانک (٪)", "نسبت نمونه (٪)", "اختلاف (واحد درصد)"])
for _, x in _c.iterrows():
    ym = str(int(x.year_month)); ws2.append([ym[4:] + "/" + ym[:4], round(x.bank_ratio, 2), round(x.sample_ratio, 2), round(x["diff"], 2)])
ws2.append([])
ws2.append(["تعداد ماه", len(_c)])
ws2.append(["همبستگی", round(_c.bank_ratio.corr(_c.sample_ratio), 3)])
ws2.append(["میانگین قدرمطلق اختلاف", round(_c["diff"].abs().mean(), 3)])
ws2.append(["ماه‌های با اختلاف کمتر از یک واحد", int((_c["diff"].abs() < 1).sum())])
for col in "ABCD":
    ws2.column_dimensions[col].width = 28
wb.save(OUT)
print("WROTE", OUT, "sheets:", order, "tables:", total)

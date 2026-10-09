"""مقایسه‌ی نمونه‌ی تحلیلی با توزیع کل مشتریان بانک (فایل RFM بانک، مقطع ۱۴۰۴/۰۹/۳۰) و آزمون
سازگاری نمونه با معیارهای پالایش طراحی نمونه‌گیری.

خروجی‌ها (data/processed/phase1_market_overview):
- loan_band_sample_vs_bank.csv    : تعداد گیرنده در طبقات میانگین مانده‌ی تسهیلات (۱۴۰۴/۰۱–۱۴۰۴/۰۹)؛ نمونه در برابر کل بانک
- deposit_band_sample_vs_bank.csv : تعداد مشتری در طبقات میانگین مانده‌ی سپرده (۱۴۰۴/۰۱–۱۴۰۴/۰۹)
- tenure_sample_vs_bank.csv       : توزیع قدمت رابطه
- sampling_filter_compliance.csv  : سهم مشتریان نمونه که هر معیار پالایش را احراز می‌کنند
- borrower_activity_primary.csv   : وضعیت مانده‌ی ۱۴۰۴ برای گیرندگان اصلی تسهیلات

فقط یک Engine (SQL Server در حالت تک‌کاربره)؛ اجرا: python -u مستقیم.
"""

import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from tt_bank.db.connection import get_engine

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "processed" / "phase1_market_overview"
RFM = ROOT / "docs" / "Sampling Method" / "RFM.xlsx"


def read_bank_bands(sheet: str, last_key: str) -> pd.DataFrame:
    """جدول‌های سالانه‌ی RFM: ردیف‌های (طبقه، سال، حقیقی، حقوقی)؛ فقط سال مورد نظر برگردانده می‌شود."""
    ws = openpyxl.load_workbook(RFM, data_only=True)[sheet]
    rows, year = [], None
    for r in ws.iter_rows(values_only=True):
        r = [c for c in r if c is not None] if sheet == "Monetary-Deposite" else list(r)
        vals = [v for v in r]
        if sheet == "Monetary-Deposite":
            label, y, nat, leg = vals[0], vals[1] if len(vals) > 3 else None, vals[-2], vals[-1]
        else:
            label, y, nat, leg = vals[0], vals[1], vals[2], vals[3]
        if y is not None:
            year = str(y)
        if isinstance(label, str) and isinstance(nat, (int, float)) and "جمع" not in label and year == last_key:
            rows.append((label.strip(), nat, leg))
    return pd.DataFrame(rows, columns=["band", "bank_natural", "bank_legal"])


def main() -> None:
    engine = get_engine()

    # ---------- تسهیلات: طبقات میانگین مانده‌ی مشتری در ۱۴۰۴/۰۱ تا ۱۴۰۴/۰۹ ----------
    # مانده‌ی هر قرارداد به «مشتری اصلی» همان قرارداد نسبت داده می‌شود (Loan_Contract_Info.MAIN_CUSTOMER_ID؛
    # برای قراردادهای فاقد رکورد قرارداد و دارای یک طرف، همان طرف). مانده‌ی ضامن‌ها به حساب نمی‌آید.
    primary = set(pd.read_sql("SELECT DISTINCT ID AS customer_id FROM dbo.Loan_Grant_Customer", engine)["customer_id"])
    print("primary list", len(primary), flush=True)
    info = pd.read_sql(
        "SELECT DISTINCT CAST(LOANFILE_ID AS numeric(38,0)) AS loanfile_id, CAST(LTRIM(RTRIM(MAIN_CUSTOMER_ID)) AS numeric(38,0)) AS main_id "
        "FROM dbo.Loan_Contract_Info WHERE ISNUMERIC(LTRIM(RTRIM(MAIN_CUSTOMER_ID))) = 1", engine).drop_duplicates("loanfile_id")
    edges = [0, 5e8, 2e9, 1e10, 5e10, 2e11, np.inf]
    loan_labels = ["کمتر از ۵۰ میلیون تومان", "۵۰ تا ۲۰۰ میلیون تومان", "۲۰۰ میلیون تا ۱ میلیارد تومان",
                   "۱ تا ۵ میلیارد تومان", "۵ تا ۲۰ میلیارد تومان", "بیش از ۲۰ میلیارد تومان"]
    bank_loan = read_bank_bands("Monetary_Loan", "1404/09/30")
    bank_loan["band"] = loan_labels
    frames = []
    for ym in range(140401, 140410):
        df = pd.read_sql(
            f"SELECT customer_id, loanfile_id, MAX(Total_Remain) AS bal FROM dbo.Loan_Remain "
            f"WHERE Transaction_Date BETWEEN {ym * 100 + 1} AND {ym * 100 + 31} GROUP BY customer_id, loanfile_id", engine)
        c = df.groupby("loanfile_id").agg(bal=("bal", "max"), nparties=("customer_id", "nunique"),
                                          first_party=("customer_id", "min")).reset_index().merge(info, on="loanfile_id", how="left")
        c["owner"] = c.main_id
        single = c.owner.isna() & (c.nparties == 1)
        c.loc[single, "owner"] = c.loc[single, "first_party"]
        c["ym"] = ym
        frames.append(c.dropna(subset=["owner"])[["owner", "ym", "bal"]])
        print("loan month", ym, flush=True)
    allm = pd.concat(frames)
    per = allm.groupby(["owner", "ym"]).bal.sum().unstack(fill_value=0).sum(axis=1) / 9.0
    per = per[per > 0]
    res = bank_loan.copy()
    res["db_customers"] = pd.cut(per, edges, right=False, labels=loan_labels).value_counts().reindex(loan_labels).values
    res["bank_total"] = res.bank_natural + res.bank_legal
    res["sample_to_bank_ratio"] = res.db_customers / res.bank_total
    res.to_csv(OUT / "loan_band_sample_vs_bank.csv", index=False, encoding="utf-8-sig")
    print(res.to_string(), flush=True)
    if "--loans-only" in sys.argv:
        return

    # ---------- گیرندگان اصلی: وضعیت مانده در ۱۴۰۴ ----------
    act = pd.read_sql("SELECT customer_id, is_active_loan_customer, max_total_remain_1404 FROM dbo.vw_Active_Loan_Customer_Flags", engine)
    act["is_primary"] = act.customer_id.isin(primary)
    summ = pd.DataFrame({
        "population": ["همه‌ی طرفین", "گیرنده‌ی اصلی"],
        "customers_with_1404_record": [len(act), int(act.is_primary.sum())],
        "with_positive_balance_1404": [int(act.is_active_loan_customer.sum()),
                                       int(act.loc[act.is_primary, "is_active_loan_customer"].sum())],
    })
    summ["share"] = summ.with_positive_balance_1404 / summ.customers_with_1404_record
    summ["primary_total_in_list"] = [np.nan, len(primary)]
    summ.to_csv(OUT / "borrower_activity_primary.csv", index=False, encoding="utf-8-sig")
    print(summ.to_string(), flush=True)

    # ---------- سپرده ----------
    q = """
    WITH m AS (
        SELECT customer_id, year_mount,
               SUM(avr_total_amount_eq) AS bal,
               MAX(CASE WHEN sum_creditor_circul_count > 0 OR sum_debtor_circul_count > 0 THEN 1 ELSE 0 END) AS t
        FROM dbo.Deposit_Monthly_Activity GROUP BY customer_id, year_mount)
    SELECT customer_id,
           MAX(bal) AS max_bal,
           SUM(t) AS active_months,
           MAX(CASE WHEN year_mount BETWEEN 140404 AND 140409 AND t = 1 THEN 1 ELSE 0 END) AS recent_to_azar,
           MAX(CASE WHEN year_mount BETWEEN 140407 AND 140412 AND t = 1 THEN 1 ELSE 0 END) AS recent_to_esfand,
           SUM(CASE WHEN year_mount BETWEEN 140401 AND 140409 THEN bal ELSE 0 END) / 9.0 AS avg_bal_1404_azar
    FROM m GROUP BY customer_id
    """
    dep = pd.read_sql(q, engine)
    print("deposit customers", len(dep), flush=True)
    acc = pd.read_sql(
        "SELECT CAST(LTRIM(RTRIM(MAIN_CUSTOMER_ID)) AS numeric(38,0)) AS customer_id, OPEN_DATE FROM dbo.Deposit_Accoutnt_Info "
        "WHERE ISNUMERIC(LTRIM(RTRIM(MAIN_CUSTOMER_ID))) = 1", engine)
    acc = acc[acc.OPEN_DATE.astype(str).str.strip().str[:4].str.isnumeric()]
    od = acc.OPEN_DATE.astype(str).str.strip()
    acc["ym"] = od.str[:4].astype(int) * 12 + od.str[5:7].astype(int)
    first_open = acc.groupby("customer_id")["ym"].min()
    all_dep = pd.DataFrame({"customer_id": pd.read_sql(
        "SELECT DISTINCT CAST(LTRIM(RTRIM(MAIN_CUSTOMER_ID)) AS numeric(38,0)) AS customer_id FROM dbo.Deposit_Accoutnt_Info "
        "WHERE ISNUMERIC(LTRIM(RTRIM(MAIN_CUSTOMER_ID))) = 1", engine)["customer_id"]})
    d = all_dep.merge(dep, on="customer_id", how="left").fillna({"max_bal": 0, "active_months": 0, "recent_to_azar": 0,
                                                                   "recent_to_esfand": 0, "avg_bal_1404_azar": 0})
    d["tenure_azar"] = 1404 * 12 + 9 - d.customer_id.map(first_open)
    d["tenure_esfand"] = 1404 * 12 + 12 - d.customer_id.map(first_open)

    # توزیع طبقات میانگین مانده
    dep_edges = [-np.inf, 1e7, 5e7, 2e8, 1e9, 5e9, np.inf]
    dep_labels = ["کمتر از ۱۰ میلیون ریال", "۱۰ تا ۵۰ میلیون ریال", "۵۰ تا ۲۰۰ میلیون ریال",
                  "۲۰۰ میلیون تا ۱ میلیارد ریال", "۱ تا ۵ میلیارد ریال", "بیش از ۵ میلیارد ریال"]
    ws = openpyxl.load_workbook(RFM, data_only=True)["Monetary-Deposite"]
    bank_dep, year = [], None
    for r in ws.iter_rows(values_only=True):
        label, y, nat, leg = r[0], r[1], r[2], r[3]
        if y is not None:
            year = str(y)
        if isinstance(label, str) and isinstance(nat, (int, float)) and "جمع" not in label and year == "1404/09/30":
            bank_dep.append((nat, leg))
    dres = pd.DataFrame(bank_dep, columns=["bank_natural", "bank_legal"])
    dres.insert(0, "band", dep_labels)
    dres["bank_total"] = dres.bank_natural + dres.bank_legal
    dres["db_customers"] = pd.cut(d.avg_bal_1404_azar, dep_edges, right=False, labels=dep_labels).value_counts().reindex(dep_labels).values
    dres.to_csv(OUT / "deposit_band_sample_vs_bank.csv", index=False, encoding="utf-8-sig")
    print(dres.to_string(), flush=True)

    # توزیع قدمت
    t_edges = [-np.inf, 6, 9, 12, 24, 60, 120, np.inf]
    t_labels = ["کمتر از ۶ ماه", "۶ تا ۹ ماه", "۹ تا ۱۲ ماه", "۱ تا ۲ سال", "۲ تا ۵ سال", "۵ تا ۱۰ سال", "بیش از ۱۰ سال"]
    bank_t = [225254, 89720, 102569, 966008, 1782836, 3036024, 2146839]
    tres = pd.DataFrame({"band": t_labels, "bank_total": bank_t})
    tres["db_customers"] = pd.cut(d.tenure_azar.dropna(), t_edges, right=False, labels=t_labels).value_counts().reindex(t_labels).values
    tres.to_csv(OUT / "tenure_sample_vs_bank.csv", index=False, encoding="utf-8-sig")
    print(tres.to_string(), flush=True)

    # سازگاری با معیارهای پالایش
    n = len(d)
    rows = []
    for tag, rec, ten in (("مقطع آذر ۱۴۰۴", "recent_to_azar", "tenure_azar"), ("مقطع اسفند ۱۴۰۴", "recent_to_esfand", "tenure_esfand")):
        c1 = d[rec] == 1
        c2 = d.active_months >= 1
        c3 = d.max_bal >= 5_000_000
        c4 = d[ten] >= 12
        rows.append((tag, n, c1.sum(), c2.sum(), c3.sum(), c4.sum(), (c1 & c2 & c3 & c4).sum()))
    comp = pd.DataFrame(rows, columns=["evaluation", "depositors", "recency", "frequency", "monetary", "tenure", "all_four"])
    comp.to_csv(OUT / "sampling_filter_compliance.csv", index=False, encoding="utf-8-sig")
    print(comp.to_string(), flush=True)


if __name__ == "__main__":
    main()
